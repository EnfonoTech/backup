"""
ePromise → ERPNext Export — Exact ERPNext Data Import Format
=============================================================
Matches what ERPNext generates when you download a template from
Settings → Data Import → Download Template.

Rules:
  • Child table column headers:  "Field Label (Table Label)"
    e.g.  "Item (Items)",  "Quantity (Items)",  "Account (Accounting Entries)"
  • Parent fields appear ONLY on the first item row of each document.
    Continuation child rows have empty parent columns.
  • The "ID" column is the document name.  Leave blank on production if
    you want ERPNext to auto-number, otherwise it forces that exact name.

Sheets:
  Customers          — Customer (one row per record)
  Suppliers          — Supplier (one row per record)
  Sales_Invoices     — Sales Invoice + Items
  Purchase_Invoices  — Purchase Invoice + Items
  Payment_Entries    — Payment Entry + References
  Journal_Entries    — Journal Entry + Accounts

Usage:
    cd /home/gym/new-bench
    ./env/bin/python apps/backup/export_to_excel.py
"""

import os, sys, datetime
os.chdir("/home/gym/new-bench")
sys.path.insert(0, "/home/gym/new-bench/apps/frappe")
sys.path.insert(0, "/home/gym/new-bench/apps/erpnext")
sys.path.insert(0, "/home/gym/new-bench/apps/backup")

import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

frappe.init(site="ksa", sites_path="/home/gym/new-bench/sites")
frappe.connect()
frappe.set_user("Administrator")

# Load unified item code map: ite_code → unified_code
# The current item_code in ERPNext SI/PI items IS the ite_code.
# For ~170 items, unified_code differs — use it as the authoritative ERPNext item_code.
from backup.epromise_migration.utils.unified_code_map import load_unified_map, get_erp_item_code
_unified_map = load_unified_map()

# Production warehouse (auto-created by ERPNext on company setup)
PROD_WAREHOUSE = "Stores - SFTB"

# ePromise cost centre codes → ERPNext cost center names
_CC_MAP = {
    "0001": "0001 - SFTB",
    "0003": "0003 - SFTB",
}


def _resolve_cost_center(raw_cc):
    """Map ePromise branch code to ERPNext cost center name."""
    if not raw_cc:
        return "Main - SFTB"
    return _CC_MAP.get(str(raw_cc).strip(), raw_cc)


def _resolve_item_code(raw_code):
    """Map a staging ERPNext item_code (= ite_code) to unified_code for production."""
    if not raw_code:
        return raw_code
    return get_erp_item_code(str(raw_code), _unified_map)

TODAY    = datetime.date.today().strftime("%Y-%m-%d")
OUT_FILE = f"/home/gym/new-bench/apps/backup/epromise_export_SFTB_{TODAY}.xlsx"

wb = Workbook()
wb.remove(wb.active)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
ALT_FILL    = PatternFill("solid", fgColor="EBF3FB")
NORM_FONT   = Font(size=10)
THIN        = Side(style="thin", color="BDD7EE")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_sheet(title, headers, data_rows):
    """Write rows to a new sheet.  data_rows is a list of lists (already ordered)."""
    ws = wb.create_sheet(title=title[:31])

    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")

    for ri, row in enumerate(data_rows, 2):
        fill = ALT_FILL if ri % 2 == 0 else None
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=(val if val is not None else ""))
            c.font = NORM_FONT; c.border = BORDER
            if fill: c.fill = fill

    # Auto-width
    for ci, h in enumerate(headers, 1):
        sample_vals = [str(data_rows[r][ci-1] or "") for r in range(min(300, len(data_rows)))]
        w = max(len(str(h)), max((len(s) for s in sample_vals), default=0))
        ws.column_dimensions[get_column_letter(ci)].width = min(w + 3, 55)

    ws.freeze_panes = "A2"
    print(f"  {title:<30} — {len(data_rows):,} rows")
    return ws


def q(sql):
    return frappe.db.sql(sql, as_dict=True)


EMPTY = ""   # sentinel for blank parent cell on continuation rows

print(f"\nExporting for ERPNext Data Import — Steel Force Trading Bahrain")
print(f"Output: {OUT_FILE}\n")


# ══════════════════════════════════════════════════════════════════════════════
# Customers  (simple — one row per record, no child table)
# ══════════════════════════════════════════════════════════════════════════════
headers = ["ID", "Customer Name", "Customer Type", "Customer Group",
           "Territory", "Tax Id", "Payment Terms", "Disabled", "ePromise Acc Code"]

rows_raw = q("""
    SELECT name AS ID, customer_name, customer_type, customer_group,
           territory, tax_id, payment_terms, disabled, epromise_acc_code
    FROM `tabCustomer`
    WHERE epromise_acc_code IS NOT NULL AND epromise_acc_code != ''
    ORDER BY epromise_acc_code
""")

data = [[r.ID, r.customer_name, r.customer_type, r.customer_group,
         r.territory, r.tax_id, r.payment_terms, r.disabled, r.epromise_acc_code]
        for r in rows_raw]
make_sheet("Customers", headers, data)


# ══════════════════════════════════════════════════════════════════════════════
# Suppliers  (simple — one row per record)
# ══════════════════════════════════════════════════════════════════════════════
headers = ["ID", "Supplier Name", "Supplier Type", "Supplier Group",
           "Tax Id", "Payment Terms", "Disabled", "ePromise Acc Code"]

rows_raw = q("""
    SELECT name AS ID, supplier_name, supplier_type, supplier_group,
           tax_id, payment_terms, disabled, epromise_acc_code
    FROM `tabSupplier`
    WHERE epromise_acc_code IS NOT NULL AND epromise_acc_code != ''
    ORDER BY epromise_acc_code
""")

data = [[r.ID, r.supplier_name, r.supplier_type, r.supplier_group,
         r.tax_id, r.payment_terms, r.disabled, r.epromise_acc_code]
        for r in rows_raw]
make_sheet("Suppliers", headers, data)


# ══════════════════════════════════════════════════════════════════════════════
# Sales Invoices
# Parent fields → first item row only; subsequent items → blank parent cols.
# Child header format: "Field Label (Items)"
# ══════════════════════════════════════════════════════════════════════════════
P_SI = ["ID", "Series", "Customer", "Date", "Payment Due Date",
        "Currency", "Exchange Rate", "Is Return (Credit Note)", "Return Against",
        "Debit To", "Company", "ePromise VR No", "ePromise TRC Code", "Remarks"]

C_SI = ["Item (Items)", "Item Name (Items)", "Quantity (Items)", "UOM (Items)",
        "Rate (Items)", "Amount (Items)", "Warehouse (Items)", "Income Account (Items)", "Cost Center (Items)"]

headers = P_SI + C_SI

# Fetch all items joined to their parent invoice
rows_raw = q("""
    SELECT
        si.name, si.naming_series, si.customer, si.posting_date, si.due_date,
        si.currency, si.conversion_rate, si.is_return, si.return_against,
        si.debit_to, si.company, si.epromise_vr_no, si.epromise_trc_code, si.remarks,
        sii.item_code, sii.item_name, sii.qty, sii.uom,
        sii.rate, sii.amount, sii.income_account, sii.cost_center,
        sii.idx
    FROM `tabSales Invoice` si
    LEFT JOIN `tabSales Invoice Item` sii ON sii.parent = si.name
    WHERE si.epromise_vr_no IS NOT NULL AND si.epromise_vr_no != ''
      AND si.docstatus != 2
    ORDER BY si.epromise_trc_code, si.epromise_vr_no, sii.idx
""")

data = []
prev_id = None
for r in rows_raw:
    is_first = (r.name != prev_id)
    prev_id  = r.name

    parent_vals = [r.name, r.naming_series, r.customer, r.posting_date, r.due_date,
                   r.currency, r.conversion_rate, r.is_return, r.return_against,
                   r.debit_to, r.company, r.epromise_vr_no, r.epromise_trc_code, r.remarks] \
                  if is_first else [""] * len(P_SI)

    child_vals  = [_resolve_item_code(r.item_code), r.item_name, r.qty, r.uom,
                   r.rate, r.amount, PROD_WAREHOUSE, r.income_account, _resolve_cost_center(r.cost_center)]

    data.append(parent_vals + child_vals)

make_sheet("Sales_Invoices", headers, data)


# ══════════════════════════════════════════════════════════════════════════════
# Purchase Invoices
# ══════════════════════════════════════════════════════════════════════════════
P_PI = ["ID", "Series", "Supplier", "Date", "Supplier Invoice No",
        "Supplier Invoice Date", "Due Date", "Currency", "Exchange Rate",
        "Is Return (Debit Note)", "Return Against Purchase Invoice",
        "Credit To", "Company", "ePromise VR No", "ePromise TRC Code", "Remarks"]

C_PI = ["Item (Items)", "Item Name (Items)", "Accepted Qty (Items)", "UOM (Items)",
        "Rate (Items)", "Amount (Items)", "Warehouse (Items)", "Expense Head (Items)", "Cost Center (Items)"]

headers = P_PI + C_PI

rows_raw = q("""
    SELECT
        pi.name, pi.naming_series, pi.supplier, pi.posting_date,
        pi.bill_no, pi.bill_date, pi.due_date,
        pi.currency, pi.conversion_rate, pi.is_return, pi.return_against,
        pi.credit_to, pi.company, pi.epromise_vr_no, pi.epromise_trc_code, pi.remarks,
        pii.item_code, pii.item_name, pii.qty, pii.uom,
        pii.rate, pii.amount, pii.expense_account, pii.cost_center,
        pii.idx
    FROM `tabPurchase Invoice` pi
    LEFT JOIN `tabPurchase Invoice Item` pii ON pii.parent = pi.name
    WHERE pi.epromise_vr_no IS NOT NULL AND pi.epromise_vr_no != ''
      AND pi.docstatus != 2
    ORDER BY pi.epromise_trc_code, pi.epromise_vr_no, pii.idx
""")

data = []
prev_id = None
for r in rows_raw:
    is_first = (r.name != prev_id)
    prev_id  = r.name

    parent_vals = [r.name, r.naming_series, r.supplier, r.posting_date,
                   r.bill_no, r.bill_date, r.due_date,
                   r.currency, r.conversion_rate, r.is_return, r.return_against,
                   r.credit_to, r.company, r.epromise_vr_no, r.epromise_trc_code, r.remarks] \
                  if is_first else [""] * len(P_PI)

    child_vals  = [_resolve_item_code(r.item_code), r.item_name, r.qty, r.uom,
                   r.rate, r.amount, PROD_WAREHOUSE, r.expense_account, _resolve_cost_center(r.cost_center)]

    data.append(parent_vals + child_vals)

make_sheet("Purchase_Invoices", headers, data)


# ══════════════════════════════════════════════════════════════════════════════
# Payment Entries
# Child header: "Field Label (Payment References)"
# ══════════════════════════════════════════════════════════════════════════════
P_PE = ["ID", "Series", "Payment Type", "Party Type", "Party",
        "Posting Date", "Paid Amount", "Received Amount",
        "Account Paid From", "Account Paid To", "Mode of Payment",
        "Cheque/Reference No", "Cheque/Reference Date",
        "Company", "ePromise VR No", "Remarks"]

C_PE = ["Type (Payment References)", "Name (Payment References)",
        "Allocated (Payment References)", "Due Date (Payment References)"]

headers = P_PE + C_PE

rows_raw = q("""
    SELECT
        pe.name, pe.naming_series, pe.payment_type, pe.party_type, pe.party,
        pe.posting_date, pe.paid_amount, pe.received_amount,
        pe.paid_from, pe.paid_to, pe.mode_of_payment,
        pe.reference_no, pe.reference_date,
        pe.company, pe.epromise_vr_no, pe.remarks,
        per.reference_doctype, per.reference_name, per.allocated_amount, per.due_date,
        per.idx
    FROM `tabPayment Entry` pe
    LEFT JOIN `tabPayment Entry Reference` per ON per.parent = pe.name
    WHERE pe.epromise_vr_no IS NOT NULL AND pe.epromise_vr_no != ''
      AND pe.docstatus != 2
    ORDER BY pe.epromise_vr_no, per.idx
""")

data = []
prev_id = None
for r in rows_raw:
    is_first = (r.name != prev_id)
    prev_id  = r.name

    parent_vals = [r.name, r.naming_series, r.payment_type, r.party_type, r.party,
                   r.posting_date, r.paid_amount, r.received_amount,
                   r.paid_from, r.paid_to, r.mode_of_payment,
                   r.reference_no, r.reference_date,
                   r.company, r.epromise_vr_no, r.remarks] \
                  if is_first else [""] * len(P_PE)

    child_vals  = [r.reference_doctype, r.reference_name, r.allocated_amount, r.due_date]

    data.append(parent_vals + child_vals)

make_sheet("Payment_Entries", headers, data)


# ══════════════════════════════════════════════════════════════════════════════
# Journal Entries
# Child header: "Field Label (Accounting Entries)"
# ══════════════════════════════════════════════════════════════════════════════
P_JE = ["ID", "Series", "Entry Type", "Posting Date",
        "Company", "ePromise VR No", "User Remark"]

C_JE = ["Account (Accounting Entries)", "Debit (Accounting Entries)",
        "Credit (Accounting Entries)", "Party Type (Accounting Entries)",
        "Party (Accounting Entries)", "Cost Center (Accounting Entries)",
        "User Remark (Accounting Entries)"]

headers = P_JE + C_JE

rows_raw = q("""
    SELECT
        je.name, je.naming_series, je.voucher_type, je.posting_date,
        je.company, je.epromise_vr_no, je.user_remark,
        jea.account, jea.debit_in_account_currency, jea.credit_in_account_currency,
        jea.party_type, jea.party, jea.cost_center,
        jea.user_remark AS line_remark,
        jea.idx
    FROM `tabJournal Entry` je
    LEFT JOIN `tabJournal Entry Account` jea ON jea.parent = je.name
    WHERE je.epromise_vr_no IS NOT NULL AND je.epromise_vr_no != ''
      AND je.docstatus != 2
    ORDER BY je.epromise_vr_no, jea.idx
""")

data = []
prev_id = None
for r in rows_raw:
    is_first = (r.name != prev_id)
    prev_id  = r.name

    parent_vals = [r.name, r.naming_series, r.voucher_type, r.posting_date,
                   r.company, r.epromise_vr_no, r.user_remark] \
                  if is_first else [""] * len(P_JE)

    child_vals  = [r.account, r.debit_in_account_currency, r.credit_in_account_currency,
                   r.party_type, r.party, r.cost_center, r.line_remark]

    data.append(parent_vals + child_vals)

make_sheet("Journal_Entries", headers, data)


# ── Save ───────────────────────────────────────────────────────────────────────
wb.save(OUT_FILE)
print(f"""
  Saved: {OUT_FILE}

  Import order in production ERPNext:
    1. Customers        → Doctype: Customer
    2. Suppliers        → Doctype: Supplier
    3. Sales_Invoices   → Doctype: Sales Invoice
    4. Purchase_Invoices→ Doctype: Purchase Invoice
    5. Payment_Entries  → Doctype: Payment Entry
    6. Journal_Entries  → Doctype: Journal Entry

  Steps per sheet:
    Settings → Data Import → New → Insert New Records → upload sheet → Import

  NOTE: Custom fields (ePromise VR No, ePromise TRC Code, ePromise Acc Code)
        must exist on the production site before importing.
        Run the migration app setup on production first.
""")

frappe.destroy()
