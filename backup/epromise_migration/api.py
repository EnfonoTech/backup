"""
Excel Import API — import the exported ERPNext Data Import Excel into production.

Whitelisted methods (called from the epromise-import page):
  get_excel_sheets(file_url)          → {sheets, order}
  import_excel_sheet(file_url, sheet_name) → {log_name, status: "queued"}
"""

import os
import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, now_datetime

COMPANY = "Steel Force Trading Bahrain"

SHEET_ORDER = [
    "Customers",
    "Suppliers",
    "Sales_Invoices",
    "Purchase_Invoices",
    "Payment_Entries",
    "Journal_Entries",
]

_MIGRATION_TYPE = {
    "Customers":         "Excel: Customers",
    "Suppliers":         "Excel: Suppliers",
    "Sales_Invoices":    "Excel: Sales Invoices",
    "Purchase_Invoices": "Excel: Purchase Invoices",
    "Payment_Entries":   "Excel: Payment Entries",
    "Journal_Entries":   "Excel: Journal Entries",
}

# Parent-column count per sheet (child cols come after)
_N_PARENT = {
    "Customers":         9,   # no child table — all cols are parent
    "Suppliers":         8,
    "Sales_Invoices":   14,
    "Purchase_Invoices": 16,
    "Payment_Entries":  16,
    "Journal_Entries":   7,
}


# ── Public whitelisted API ────────────────────────────────────────────────────

@frappe.whitelist()
def get_excel_sheets(file_url):
    """Return sheets present in the uploaded Excel that match the expected order."""
    import openpyxl
    path = _resolve_file_path(file_url)
    wb = openpyxl.load_workbook(path, read_only=True)
    available = set(wb.sheetnames)
    wb.close()
    return {
        "sheets": [s for s in SHEET_ORDER if s in available],
        "order":  SHEET_ORDER,
    }


@frappe.whitelist()
def import_excel_sheet(file_url, sheet_name):
    """Enqueue the import of one sheet. Returns {log_name, status: 'queued'}."""
    if sheet_name not in SHEET_ORDER:
        frappe.throw(_(f"Unknown sheet: {sheet_name}"))

    migration_type = _MIGRATION_TYPE[sheet_name]

    log = frappe.get_doc({
        "doctype": "ePromise Migration Log",
        "migration_type": migration_type,
        "status": "Queued",
        "started_at": now_datetime(),
    })
    log.insert(ignore_permissions=True)
    frappe.db.commit()

    frappe.enqueue(
        "backup.epromise_migration.api._run_sheet_import",
        queue="long",
        timeout=3600,
        file_url=file_url,
        sheet_name=sheet_name,
        log_name=log.name,
    )

    return {"log_name": log.name, "status": "queued"}


# ── Background worker ─────────────────────────────────────────────────────────

def _run_sheet_import(file_url, sheet_name, log_name):
    """Background job: parse Excel sheet and insert records."""
    import openpyxl

    log = frappe.get_doc("ePromise Migration Log", log_name)
    log.status = "Running"
    log.save(ignore_permissions=True)
    frappe.db.commit()

    errors_log = []
    counts = {"total": 0, "created": 0, "skipped": 0, "errors": 0}

    try:
        path = _resolve_file_path(file_url)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")

        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if len(all_rows) < 2:
            raise ValueError(f"Sheet '{sheet_name}' has no data rows.")

        headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
        data_rows = all_rows[1:]

        n_parent = _N_PARENT[sheet_name]
        handler  = _HANDLERS[sheet_name]

        if sheet_name in ("Customers", "Suppliers"):
            # Flat (no child table)
            for row in data_rows:
                if not any(row):
                    continue
                row_dict = dict(zip(headers, row))
                counts["total"] += 1
                result, err = handler(row_dict, headers)
                if result == "created":
                    counts["created"] += 1
                elif result == "skipped":
                    counts["skipped"] += 1
                else:
                    counts["errors"] += 1
                    if err:
                        errors_log.append(err)
                if counts["total"] % 50 == 0:
                    frappe.db.commit()
        else:
            # Parent-child reconstruction
            docs = _reconstruct_docs(headers, data_rows, n_parent)
            counts["total"] = len(docs)
            for doc_data in docs:
                result, err = handler(doc_data)
                if result == "created":
                    counts["created"] += 1
                elif result == "skipped":
                    counts["skipped"] += 1
                else:
                    counts["errors"] += 1
                    if err:
                        errors_log.append(err)
                if (counts["created"] + counts["errors"] + counts["skipped"]) % 50 == 0:
                    frappe.db.commit()

        frappe.db.commit()

        log.reload()
        log.status = "Completed"
        log.total_records  = counts["total"]
        log.success_count  = counts["created"]
        log.skipped_count  = counts["skipped"]
        log.error_count    = counts["errors"]
        log.completed_at   = now_datetime()
        log.log_details    = "\n".join(errors_log[-200:]) if errors_log else ""
        log.save(ignore_permissions=True)
        frappe.db.commit()

    except Exception as exc:
        frappe.log_error(frappe.get_traceback(), "Excel Import Error")
        log.reload()
        log.status       = "Failed"
        log.log_details  = str(exc) + "\n" + frappe.get_traceback()
        log.completed_at = now_datetime()
        log.save(ignore_permissions=True)
        frappe.db.commit()


# ── Sheet handlers ────────────────────────────────────────────────────────────

def _import_customer_row(row_dict, headers):
    acc_code = _v(row_dict.get("ePromise Acc Code"))
    if not acc_code:
        return "skipped", None

    if frappe.db.exists("Customer", {"epromise_acc_code": acc_code}):
        return "skipped", None

    try:
        name = _v(row_dict.get("Customer Name")) or acc_code
        doc = frappe.get_doc({
            "doctype":        "Customer",
            "customer_name":  name,
            "customer_type":  _v(row_dict.get("Customer Type")) or "Company",
            "customer_group": _v(row_dict.get("Customer Group")) or "All Customer Groups",
            "territory":      _v(row_dict.get("Territory")) or "All Territories",
            "tax_id":         _v(row_dict.get("Tax Id")),
            "payment_terms":  _v(row_dict.get("Payment Terms")),
            "disabled":       cint(row_dict.get("Disabled")),
            "epromise_acc_code": acc_code,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        frappe.db.commit()
        return "created", None
    except Exception as e:
        return "error", f"Customer {acc_code}: {e}"


def _import_supplier_row(row_dict, headers):
    acc_code = _v(row_dict.get("ePromise Acc Code"))
    if not acc_code:
        return "skipped", None

    if frappe.db.exists("Supplier", {"epromise_acc_code": acc_code}):
        return "skipped", None

    try:
        name = _v(row_dict.get("Supplier Name")) or acc_code
        doc = frappe.get_doc({
            "doctype":        "Supplier",
            "supplier_name":  name,
            "supplier_type":  _v(row_dict.get("Supplier Type")) or "Company",
            "supplier_group": _v(row_dict.get("Supplier Group")) or "All Supplier Groups",
            "tax_id":         _v(row_dict.get("Tax Id")),
            "payment_terms":  _v(row_dict.get("Payment Terms")),
            "disabled":       cint(row_dict.get("Disabled")),
            "epromise_acc_code": acc_code,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        frappe.db.commit()
        return "created", None
    except Exception as e:
        return "error", f"Supplier {acc_code}: {e}"


def _import_sales_invoice(doc_data):
    p = doc_data["parent"]
    vr_no = _v(p.get("ePromise VR No"))
    if not vr_no:
        return "skipped", None
    if frappe.db.exists("Sales Invoice", {"epromise_vr_no": vr_no, "docstatus": ("!=", 2)}):
        return "skipped", None

    items = []
    for c in doc_data["children"]:
        item_code = _v(c.get("Item (Items)"))
        if not item_code:
            continue
        items.append({
            "item_code":      item_code,
            "item_name":      _v(c.get("Item Name (Items)")),
            "qty":            flt(c.get("Quantity (Items)")),
            "uom":            _v(c.get("UOM (Items)")) or "Nos",
            "rate":           flt(c.get("Rate (Items)")),
            "amount":         flt(c.get("Amount (Items)")),
            "income_account": _v(c.get("Income Account (Items)")),
            "cost_center":    _v(c.get("Cost Center (Items)")),
        })

    if not items:
        return "skipped", None

    # Clear return_against if the referenced doc won't exist in production
    return_against = _v(p.get("Return Against"))
    if return_against and not frappe.db.exists("Sales Invoice", return_against):
        return_against = None

    try:
        doc = frappe.get_doc({
            "doctype":          "Sales Invoice",
            "naming_series":    _v(p.get("Series")),
            "customer":         _v(p.get("Customer")),
            "posting_date":     getdate(_v(p.get("Date"))),
            "due_date":         getdate(_v(p.get("Payment Due Date"))),
            "currency":         _v(p.get("Currency")) or "BHD",
            "conversion_rate":  flt(p.get("Exchange Rate") or 1),
            "is_return":        cint(p.get("Is Return (Credit Note)")),
            "return_against":   return_against,
            "debit_to":         _v(p.get("Debit To")),
            "company":          _v(p.get("Company")) or COMPANY,
            "epromise_vr_no":   vr_no,
            "epromise_trc_code": _v(p.get("ePromise TRC Code")),
            "remarks":          _v(p.get("Remarks")),
            "items":            items,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        doc.submit()
        return "created", None
    except Exception as e:
        return "error", f"SI {vr_no}: {e}"


def _import_purchase_invoice(doc_data):
    p = doc_data["parent"]
    vr_no = _v(p.get("ePromise VR No"))
    if not vr_no:
        return "skipped", None
    if frappe.db.exists("Purchase Invoice", {"epromise_vr_no": vr_no, "docstatus": ("!=", 2)}):
        return "skipped", None

    items = []
    for c in doc_data["children"]:
        item_code = _v(c.get("Item (Items)"))
        if not item_code:
            continue
        items.append({
            "item_code":       item_code,
            "item_name":       _v(c.get("Item Name (Items)")),
            "qty":             flt(c.get("Accepted Qty (Items)")),
            "uom":             _v(c.get("UOM (Items)")) or "Nos",
            "rate":            flt(c.get("Rate (Items)")),
            "amount":          flt(c.get("Amount (Items)")),
            "expense_account": _v(c.get("Expense Head (Items)")),
            "cost_center":     _v(c.get("Cost Center (Items)")),
        })

    if not items:
        return "skipped", None

    return_against = _v(p.get("Return Against Purchase Invoice"))
    if return_against and not frappe.db.exists("Purchase Invoice", return_against):
        return_against = None

    try:
        doc = frappe.get_doc({
            "doctype":          "Purchase Invoice",
            "naming_series":    _v(p.get("Series")),
            "supplier":         _v(p.get("Supplier")),
            "posting_date":     getdate(_v(p.get("Date"))),
            "bill_no":          _v(p.get("Supplier Invoice No")),
            "bill_date":        getdate(_v(p.get("Supplier Invoice Date"))) if _v(p.get("Supplier Invoice Date")) else None,
            "due_date":         getdate(_v(p.get("Due Date"))) if _v(p.get("Due Date")) else None,
            "currency":         _v(p.get("Currency")) or "BHD",
            "conversion_rate":  flt(p.get("Exchange Rate") or 1),
            "is_return":        cint(p.get("Is Return (Debit Note)")),
            "return_against":   return_against,
            "credit_to":        _v(p.get("Credit To")),
            "company":          _v(p.get("Company")) or COMPANY,
            "epromise_vr_no":   vr_no,
            "epromise_trc_code": _v(p.get("ePromise TRC Code")),
            "remarks":          _v(p.get("Remarks")),
            "items":            items,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        doc.submit()
        return "created", None
    except Exception as e:
        return "error", f"PI {vr_no}: {e}"


def _import_payment_entry(doc_data):
    p = doc_data["parent"]
    vr_no = _v(p.get("ePromise VR No"))
    if not vr_no:
        return "skipped", None
    if frappe.db.exists("Payment Entry", {"epromise_vr_no": vr_no, "docstatus": ("!=", 2)}):
        return "skipped", None

    # Rebuild references only for refs that exist in production
    references = []
    for c in doc_data["children"]:
        ref_doctype = _v(c.get("Type (Payment References)"))
        ref_name    = _v(c.get("Name (Payment References)"))
        if ref_doctype and ref_name and frappe.db.exists(ref_doctype, ref_name):
            references.append({
                "reference_doctype": ref_doctype,
                "reference_name":    ref_name,
                "allocated_amount":  flt(c.get("Allocated (Payment References)")),
                "due_date":          getdate(_v(c.get("Due Date (Payment References)"))) if _v(c.get("Due Date (Payment References)")) else None,
            })

    try:
        doc = frappe.get_doc({
            "doctype":          "Payment Entry",
            "naming_series":    _v(p.get("Series")),
            "payment_type":     _v(p.get("Payment Type")),
            "party_type":       _v(p.get("Party Type")),
            "party":            _v(p.get("Party")),
            "posting_date":     getdate(_v(p.get("Posting Date"))),
            "paid_amount":      flt(p.get("Paid Amount")),
            "received_amount":  flt(p.get("Received Amount")),
            "paid_from":        _v(p.get("Account Paid From")),
            "paid_to":          _v(p.get("Account Paid To")),
            "mode_of_payment":  _v(p.get("Mode of Payment")),
            "reference_no":     _v(p.get("Cheque/Reference No")),
            "reference_date":   getdate(_v(p.get("Cheque/Reference Date"))) if _v(p.get("Cheque/Reference Date")) else None,
            "company":          _v(p.get("Company")) or COMPANY,
            "epromise_vr_no":   vr_no,
            "remarks":          _v(p.get("Remarks")),
            "references":       references,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        doc.submit()
        return "created", None
    except Exception as e:
        return "error", f"PE {vr_no}: {e}"


def _import_journal_entry(doc_data):
    p = doc_data["parent"]
    vr_no = _v(p.get("ePromise VR No"))
    if not vr_no:
        return "skipped", None
    if frappe.db.exists("Journal Entry", {"epromise_vr_no": vr_no, "docstatus": ("!=", 2)}):
        return "skipped", None

    accounts = []
    for c in doc_data["children"]:
        acct = _v(c.get("Account (Accounting Entries)"))
        if not acct:
            continue
        accounts.append({
            "account":                    acct,
            "debit_in_account_currency":  flt(c.get("Debit (Accounting Entries)")),
            "credit_in_account_currency": flt(c.get("Credit (Accounting Entries)")),
            "party_type":                 _v(c.get("Party Type (Accounting Entries)")),
            "party":                      _v(c.get("Party (Accounting Entries)")),
            "cost_center":                _v(c.get("Cost Center (Accounting Entries)")),
            "user_remark":                _v(c.get("User Remark (Accounting Entries)")),
        })

    if not accounts:
        return "skipped", None

    try:
        doc = frappe.get_doc({
            "doctype":        "Journal Entry",
            "naming_series":  _v(p.get("Series")),
            "voucher_type":   _v(p.get("Entry Type")),
            "posting_date":   getdate(_v(p.get("Posting Date"))),
            "company":        _v(p.get("Company")) or COMPANY,
            "epromise_vr_no": vr_no,
            "user_remark":    _v(p.get("User Remark")),
            "accounts":       accounts,
        })
        doc.flags.ignore_permissions = True
        doc.flags.ignore_mandatory   = True
        doc.insert()
        doc.submit()
        return "created", None
    except Exception as e:
        return "error", f"JE {vr_no}: {e}"


_HANDLERS = {
    "Customers":         _import_customer_row,
    "Suppliers":         _import_supplier_row,
    "Sales_Invoices":    _import_sales_invoice,
    "Purchase_Invoices": _import_purchase_invoice,
    "Payment_Entries":   _import_payment_entry,
    "Journal_Entries":   _import_journal_entry,
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _v(val):
    """Coerce cell value to stripped string or None."""
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def _resolve_file_path(file_url):
    """Convert a Frappe file_url to an absolute filesystem path."""
    site_path = frappe.get_site_path()
    if file_url.startswith("/files/"):
        return os.path.join(site_path, "public", "files", file_url[7:])
    elif file_url.startswith("/private/files/"):
        return os.path.join(site_path, "private", "files", file_url[15:])
    else:
        # Try File doctype lookup
        try:
            f = frappe.get_doc("File", {"file_url": file_url})
            return f.get_full_path()
        except Exception:
            return file_url


def _reconstruct_docs(headers, rows, n_parent_cols):
    """
    Reconstruct parent+children dicts from flat ERPNext Data Import rows.
    A row with non-empty col[0] (ID) starts a new parent document.
    """
    parent_headers = headers[:n_parent_cols]
    child_headers  = headers[n_parent_cols:]

    docs = []
    current = None

    for row in rows:
        if not any(r is not None and str(r).strip() for r in row):
            continue  # skip blank rows

        id_val = str(row[0]).strip() if row[0] is not None else ""

        if id_val:
            # New parent row
            if current is not None:
                docs.append(current)
            parent_dict = dict(zip(parent_headers, row[:n_parent_cols]))
            current = {"parent": parent_dict, "children": []}
            # Include child part of this same row if present
            child_row = row[n_parent_cols:]
            if any(c is not None and str(c).strip() for c in child_row):
                current["children"].append(dict(zip(child_headers, child_row)))
        else:
            # Continuation child row
            if current is not None:
                child_row = row[n_parent_cols:]
                if any(c is not None and str(c).strip() for c in child_row):
                    current["children"].append(dict(zip(child_headers, child_row)))

    if current is not None:
        docs.append(current)

    return docs
